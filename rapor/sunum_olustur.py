# -*- coding: utf-8 -*-
"""
Bitirme Projesi SUNUMU uretici (PowerPoint .pptx).
Calistirmak icin:  python rapor/sunum_olustur.py
Cikti:             rapor/Bitirme_Sunum.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- Renkler (raporla ayni kurumsal palet) ---
ANA = RGBColor(0x1F, 0x3A, 0x5F)     # koyu lacivert
VURGU = RGBColor(0x2E, 0x6D, 0xA4)   # mavi
ACIK = RGBColor(0xEC, 0xF1, 0xF7)    # cok acik mavi (arka plan)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)
GRI = RGBColor(0x55, 0x55, 0x55)
YESIL = RGBColor(0x2E, 0x8B, 0x57)

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def arka_plan(slide, renk):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = renk


def kutu(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h).text_frame


def yazi(tf, metin, boyut, renk, bold=False, hiza=PP_ALIGN.LEFT, ekle=False, italic=False):
    par = tf.paragraphs[0] if not ekle else tf.add_paragraph()
    par.alignment = hiza
    run = par.add_run()
    run.text = metin
    run.font.size = Pt(boyut)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = renk
    run.font.name = "Calibri"
    return par


def serit(slide):
    """Ust renk seridi (baslik slaytlari haric icerik slaytlarinda)."""
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(1.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = ANA
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def icerik_slayt(baslik, maddeler, alt=None):
    """Baslikli, maddeli standart icerik slayti. maddeler: [(metin, seviye), ...] veya [metin,...]"""
    s = prs.slides.add_slide(BLANK)
    arka_plan(s, BEYAZ)
    serit(s)
    # baslik
    tf = kutu(s, Inches(0.6), Inches(0.18), Inches(12.1), Inches(0.8))
    yazi(tf, baslik, 30, BEYAZ, bold=True)
    # govde
    body = kutu(s, Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.6))
    body.word_wrap = True
    ilk = True
    for m in maddeler:
        if isinstance(m, tuple):
            metin, sev = m
        else:
            metin, sev = m, 0
        par = body.paragraphs[0] if ilk else body.add_paragraph()
        ilk = False
        par.alignment = PP_ALIGN.LEFT
        par.space_after = Pt(10)
        par.level = sev
        run = par.add_run()
        isaret = "•  " if sev == 0 else "–  "
        run.text = isaret + metin
        run.font.size = Pt(22 if sev == 0 else 18)
        run.font.color.rgb = ANA if sev == 0 else GRI
        run.font.bold = (sev == 0)
        run.font.name = "Calibri"
    if alt:
        af = kutu(s, Inches(0.8), Inches(6.85), Inches(11.7), Inches(0.5))
        yazi(af, alt, 13, VURGU, italic=True)
    return s


def gorsel_yer(slide, metin):
    """Ekran goruntusu icin yer tutucu kutu."""
    box = slide.shapes.add_shape(1, Inches(7.0), Inches(1.5), Inches(5.5), Inches(4.8))
    box.fill.solid(); box.fill.fore_color.rgb = ACIK
    box.line.color.rgb = VURGU; box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "[ EKRAN GÖRÜNTÜSÜ:\n" + metin + " ]"
    r.font.size = Pt(14); r.font.italic = True; r.font.color.rgb = VURGU


# ============================================================
# 1. KAPAK
# ============================================================
s = prs.slides.add_slide(BLANK)
arka_plan(s, ANA)
# mavi aksan blok
acc = s.shapes.add_shape(1, 0, Inches(5.0), SW, Inches(0.12))
acc.fill.solid(); acc.fill.fore_color.rgb = VURGU; acc.line.fill.background()
acc.shadow.inherit = False
tf = kutu(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(2.5))
yazi(tf, "Yapay Zeka Destekli", 26, RGBColor(0x9F, 0xC0, 0xE0), bold=False)
yazi(tf, "İşletme ve Cihaz Bakım Yönetim Sistemi", 40, BEYAZ, bold=True, ekle=True)
yazi(tf, "Web + Mobil + Yapay Zeka Servisi", 18, RGBColor(0xC8, 0xD8, 0xE8), italic=True, ekle=True)
tf2 = kutu(s, Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.6))
yazi(tf2, "Hazırlayan: [ ADINIZ SOYADINIZ ]   |   Öğrenci No: [ NUMARANIZ ]", 16, BEYAZ)
yazi(tf2, "Danışman: [ DANIŞMAN ÖĞRETİM ÜYESİ ]", 16, BEYAZ, ekle=True)
yazi(tf2, "Bitirme Projesi — Haziran 2026", 14, RGBColor(0xC8, 0xD8, 0xE8), italic=True, ekle=True)

# ============================================================
# 2. PROBLEM & AMAÇ
# ============================================================
icerik_slayt("Problem ve Amaç", [
    ("Problem", 0),
    ("Cihaz bakım/servis işletmeleri süreçleri kağıt, Excel ve kopuk araçlarla yürütüyor", 1),
    ("Periyodik bakım tarihleri unutuluyor → müşteri kaybı", 1),
    ("Satış, borç, tahsilat ve saha görevleri denetlenemiyor", 1),
    ("Amaç", 0),
    ("Tüm süreçleri tek sistemde toplayan, web + mobil, çok kullanıcılı ve", 1),
    ("yapay zeka ile zenginleştirilmiş bir işletme yönetim sistemi geliştirmek", 1),
], alt="Tek işletme değil; kayıt olan birden çok işletmeye verileri yalıtarak hizmet verir (multi-tenant).")

# ============================================================
# 3. ÇÖZÜM — SİSTEM NE YAPIYOR
# ============================================================
icerik_slayt("Çözüm — Sistem Neler Yapıyor?", [
    ("Müşteri, ürün, satış, bakım ve görev yönetimi tek panelde", 0),
    ("Yaklaşan bakımların otomatik e-posta ile hatırlatılması", 0),
    ("Sahadaki teknisyen için mobil uygulama: fotoğraflı iş kanıtı + navigasyon", 0),
    ("Gelir-gider, fatura (PDF), raporlama (PDF/Excel)", 0),
    ("Harita üzerinde müşteri konumları ve rota optimizasyonu", 0),
    ("5 ayrı yapay zeka destekli özellik", 0),
])

# ============================================================
# 4. TEKNOLOJİ YIĞINI
# ============================================================
icerik_slayt("Kullanılan Teknolojiler", [
    ("Backend: ASP.NET Core (.NET 10) Web API + Entity Framework Core", 0),
    ("Veritabanı: PostgreSQL", 0),
    ("Web: React 19 + Vite + TailwindCSS", 0),
    ("Mobil: React Native (Expo) — Android", 0),
    ("Yapay Zeka: Python + FastAPI mikroservisi", 0),
    ("Makine Öğrenmesi: scikit-learn (Random Forest)", 0),
    ("Büyük Dil Modeli: Google Gemini (görüntü, ses, metin)", 0),
    ("Diğer: JWT, Hangfire, MailKit, QuestPDF, Leaflet", 0),
], alt="Her bileşen kendi güçlü olduğu alanda: tablo verisi için klasik ML, görüntü/ses/metin için LLM.")

# ============================================================
# 5. MİMARİ
# ============================================================
icerik_slayt("Sistem Mimarisi", [
    ("3 istemci → tek merkezî REST API → PostgreSQL", 0),
    ("Web paneli, Android mobil uygulaması, Yapay Zeka servisi", 1),
    ("Backend N-katmanlı (Clean Architecture):", 0),
    ("Domain → Application → Infrastructure → API", 1),
    ("Controller arayüzü çağırır; somut servis Infrastructure'da (dependency inversion)", 1),
], alt="Katmanların ayrılması → test edilebilir, bakımı kolay, genişletilebilir kod.")
# mimari slaytina sema yer tutucu ekleyelim (son slayti al)
gorsel_yer(prs.slides[-1], "Mimari şeması\n(Web+Mobil+AI → API → DB)")

# ============================================================
# 6. ÇOK KİRACILI YAPI & GÜVENLİK
# ============================================================
icerik_slayt("Çok Kiracılı Yapı ve Güvenlik", [
    ("Multi-tenant: her kayıt bir IsletmeId taşır", 0),
    ("EF Core global sorgu filtresi → her sorgu otomatik işletmeye göre süzülür", 1),
    ("Bir işletme başka işletmenin verisini ASLA göremez", 1),
    ("JWT tabanlı kimlik doğrulama, parolalar hash'li", 0),
    ("Rol bazlı yetkilendirme: Yönetici / Satış Danışmanı / Teknisyen", 0),
    ("Dosya yükleme: boyut sınırı + içerik (magic-byte) doğrulaması", 0),
    ("Merkezî hata yönetimi (uygun HTTP kodu + TraceId)", 0),
])

# ============================================================
# 7. MODÜLLER (ekran goruntusu)
# ============================================================
s = icerik_slayt("Modüller", [
    ("Gösterge Paneli (Dashboard) — özet + grafikler", 0),
    ("Müşteri / Ürün / Satış yönetimi", 0),
    ("Bakım & Servis takibi", 0),
    ("Görev yönetimi (mobil iş kanıtı)", 0),
    ("Gelir-Gider / Fatura / Raporlama", 0),
    ("Harita ve konum", 0),
])
gorsel_yer(s, "Dashboard veya\nbir modül ekranı")

# ============================================================
# 8. YAPAY ZEKA ÖZELLİKLERİ (vurgu slayti)
# ============================================================
s = prs.slides.add_slide(BLANK)
arka_plan(s, ANA)
tf = kutu(s, Inches(0.7), Inches(0.4), Inches(12), Inches(1.0))
yazi(tf, "★  5 Yapay Zeka Özelliği", 34, BEYAZ, bold=True)
body = kutu(s, Inches(0.9), Inches(1.7), Inches(11.6), Inches(5.4))
body.word_wrap = True
ai = [
    ("1. Arıza Riski Tahmini", "Makine öğrenmesi (Random Forest) — cihazın arıza riskini hesaplar"),
    ("2. Fatura Okuma (OCR)", "Gemini Vision — fatura görselinden firma/tarih/tutar/kalem çıkarır"),
    ("3. Sesle Rapor", "Teknisyenin sesi → yazıya + profesyonel teknik rapora dönüştürülür"),
    ("4. Cihaz Foto Arıza Analizi", "Gemini Vision — arızalı cihaz fotoğrafından Türkçe analiz"),
    ("5. Rota Optimizasyonu", "Teknisyenin günlük durakları en kısa rotaya dizilir (TSP)"),
]
ilk = True
for bas, ack in ai:
    par = body.paragraphs[0] if ilk else body.add_paragraph(); ilk = False
    par.space_after = Pt(14)
    r1 = par.add_run(); r1.text = bas + "  —  "
    r1.font.size = Pt(22); r1.font.bold = True; r1.font.color.rgb = RGBColor(0x9F, 0xC0, 0xE0)
    r2 = par.add_run(); r2.text = ack
    r2.font.size = Pt(18); r2.font.color.rgb = BEYAZ

# ============================================================
# 9. DEMO
# ============================================================
s = prs.slides.add_slide(BLANK)
arka_plan(s, VURGU)
tf = kutu(s, Inches(1), Inches(2.7), Inches(11.3), Inches(2))
yazi(tf, "▶  CANLI DEMO", 44, BEYAZ, bold=True, hiza=PP_ALIGN.CENTER)
yazi(tf, "Giriş · Dashboard · Bakım & Günün Rotası · AI özellikleri", 20,
     RGBColor(0xE0, 0xEC, 0xF7), hiza=PP_ALIGN.CENTER, ekle=True)

# ============================================================
# 10. SONUÇ & GELECEK
# ============================================================
icerik_slayt("Sonuç ve Gelecek Çalışmalar", [
    ("Sonuç", 0),
    ("Web + mobil + AI bileşenli, çok kullanıcılı bütünleşik sistem geliştirildi", 1),
    ("Katmanlı mimari + çok kiracılı yalıtım + rol bazlı güvenlik", 1),
    ("Gelecek Çalışmalar", 0),
    ("AI modelinin gerçek işletme verisiyle yeniden eğitilmesi", 1),
    ("EF Core Migrations'a geçiş, otomatik test ve CI/CD", 1),
    ("iOS uygulaması ve müşteri self-servis randevu portalı", 1),
])

# ============================================================
# 11. TEŞEKKÜR
# ============================================================
s = prs.slides.add_slide(BLANK)
arka_plan(s, ANA)
tf = kutu(s, Inches(1), Inches(2.9), Inches(11.3), Inches(1.8))
yazi(tf, "Teşekkür Ederim", 44, BEYAZ, bold=True, hiza=PP_ALIGN.CENTER)
yazi(tf, "Sorularınız?", 24, RGBColor(0x9F, 0xC0, 0xE0), hiza=PP_ALIGN.CENTER, ekle=True)

prs.save("rapor/Bitirme_Sunum.pptx")
print("OLUSTURULDU: rapor/Bitirme_Sunum.pptx  | Slayt sayisi:", len(prs.slides._sldIdLst))
